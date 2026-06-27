from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "World Cup Bracket.xlsx"

PEOPLE = [
    "Morgane",
    "Eileen",
    "Derek",
    "Matt M",
    "Britton",
    "Dawn",
    "Cassie",
    "Changsu",
    "Vishakha",
    "Robert",
    "Kate",
    "Tomiwa",
    "Mohamed",
    "Sergio",
    "Zarai",
    "Wojciech",
    "Cathy",
    "Mac",
    "Jeremy",
    "Tanya",
    "Jess",
    "Ryan",
    "Seeyon",
]

ROUNDS = [
    {"id": "R32", "name": "Round of 32", "shortName": "R32", "points": 1, "count": 16},
    {"id": "R16", "name": "Round of 16", "shortName": "R16", "points": 2, "count": 8},
    {"id": "QF", "name": "Quarterfinals", "shortName": "QF", "points": 4, "count": 4},
    {"id": "SF", "name": "Semifinals", "shortName": "SF", "points": 8, "count": 2},
    {"id": "FINAL", "name": "Final", "shortName": "Final", "points": 12, "count": 1},
]

COUNTRY_NAMES = {
    "ALG": "Algeria",
    "ARG": "Argentina",
    "AUS": "Australia",
    "AUT": "Austria",
    "BEL": "Belgium",
    "BIH": "Bosnia and Herzegovina",
    "BRA": "Brazil",
    "CAN": "Canada",
    "CIV": "Cote d'Ivoire",
    "COD": "DR Congo",
    "COL": "Colombia",
    "CPV": "Cape Verde",
    "CRO": "Croatia",
    "CUW": "Curacao",
    "CZE": "Czechia",
    "ECU": "Ecuador",
    "EGY": "Egypt",
    "ENG": "England",
    "ESP": "Spain",
    "FRA": "France",
    "GER": "Germany",
    "GHA": "Ghana",
    "HAI": "Haiti",
    "IRN": "Iran",
    "IRQ": "Iraq",
    "JOR": "Jordan",
    "JPN": "Japan",
    "KOR": "South Korea",
    "KSA": "Saudi Arabia",
    "MAR": "Morocco",
    "MEX": "Mexico",
    "NED": "Netherlands",
    "NOR": "Norway",
    "NZL": "New Zealand",
    "PAN": "Panama",
    "PAR": "Paraguay",
    "POR": "Portugal",
    "QAT": "Qatar",
    "RSA": "South Africa",
    "SCO": "Scotland",
    "SEN": "Senegal",
    "SUI": "Switzerland",
    "SWE": "Sweden",
    "TUN": "Tunisia",
    "TUR": "Turkey",
    "URU": "Uruguay",
    "USA": "United States",
    "UZB": "Uzbekistan",
}


def match_id(round_id: str, index: int) -> str:
    return f"{round_id}-{index:02d}"


def match_rows() -> list[list[object]]:
    rows = []
    for round_def in ROUNDS:
        for index in range(1, round_def["count"] + 1):
            rows.append([match_id(round_def["id"], index), round_def["id"], index, "", "", ""])
    return rows


def team_rows() -> list[list[str]]:
    rows = []
    for flag in sorted((ROOT / "flags").glob("*/*.png")):
        code = flag.stem.upper()
        rows.append([code, COUNTRY_NAMES.get(code, code), flag.relative_to(ROOT).as_posix()])
    return sorted(rows, key=lambda row: row[0])


def write_table(sheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="171717")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(bottom=thin)

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 3, 12), 32)


def main() -> None:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    actual = workbook.create_sheet("Actual Bracket")
    predictions = workbook.create_sheet("Bracket Predictions")
    dark_horse = workbook.create_sheet("Dark Horse Picks")
    teams = workbook.create_sheet("Team Directory")

    readme_rows = [
        ["Sheet", "Purpose", "Editable?", "Notes"],
        ["Actual Bracket", "Official knockout match teams and winners", "Yes", "Fill Team 1, Team 2, and Winner with team codes after each result."],
        ["Bracket Predictions", "One row per participant per match", "Yes", "Fill Winner with the submitted pick for that match."],
        ["Dark Horse Picks", "Separate side-quest tracking", "Yes", "Points are separate from bracket scoring."],
        ["Team Directory", "Team code, display name, and flag path", "Yes", "Add teams here if the flag asset is added later."],
        ["Scoring", "R32=1, R16=2, QF=4, SF=8, Champion=12", "No", "The website calculates totals from these rules."],
        ["Export", "Run python3 tools/export_bracket.py", "No", "This writes bracket-data.js for the browser."],
    ]
    readme.merge_cells("A1:D1")
    readme["A1"] = "World Cup Bracket Tracker Source"
    readme["A1"].fill = PatternFill("solid", fgColor="0F2F2B")
    readme["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    readme.append([])
    for row in readme_rows:
        readme.append(row)
    style_sheet(readme)

    write_table(actual, ["Match ID", "Round", "Match No", "Team 1", "Team 2", "Winner"], match_rows())

    prediction_rows = []
    for person in PEOPLE:
        for row in match_rows():
            prediction_rows.append([person, row[0], ""])
    write_table(predictions, ["Person", "Match ID", "Winner"], prediction_rows)

    write_table(dark_horse, ["Person", "Team", "Result", "Points"], [[person, "", "", 0] for person in PEOPLE])
    write_table(teams, ["Code", "Name", "Flag"], team_rows())

    workbook.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
